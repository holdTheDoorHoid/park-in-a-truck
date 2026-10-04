#!/usr/bin/env python3
"""Trace and test-recalculate Park in a Truck's cost-estimator spreadsheet.

The Dream workbook (p.18) links a Google Sheet, exported here as
source/linked/04_Dream_WORKBOOK_p18_DOWNLOAD_COST_ESTIMATOR__*.xlsx. The site
ports it to TypeScript (src/lib/cost/model.ts). This script produces the
evidence that the port is faithful:

  dump      write every non-empty cell (formula, cached value, fill colour) of
            every sheet to source/work/cost/dump_<sheet>.txt  (for reading)
  verify    recalculate the UNTOUCHED file in LibreOffice and compare every
            formula cell with the value Google Sheets cached when it exported
            the file (proves LibreOffice evaluates this sheet like Google does)
  fixtures  for each input set below: write the inputs into the orange cells
            with openpyxl, recalculate in LibreOffice, read the results back
            and save src/lib/cost/__tests__/fixtures/<name>.json

Run from the repo root:
    ~/Desktop/park-in-a-truck/source/.venv/bin/python scripts/analyze_cost_model.py verify fixtures

LibreOffice runs with a private profile (source/work/cost/lo-profile) that sets
"recalculate on load: always", so it never trusts stale cached values. openpyxl
drops cached values when it saves, and the script also checks that the results
moved with the inputs (park area and perimeter cells), so a skipped
recalculation cannot slip through.

Every function the sheet uses (SUM, SUMIF, COUNTIF, ROUNDUP) exists in
LibreOffice; none of the formulas is Google-only.
"""

from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl

HOME = Path.home()
SOURCE = HOME / 'Desktop/park-in-a-truck/source'
XLSX = SOURCE / 'linked/04_Dream_WORKBOOK_p18_DOWNLOAD_COST_ESTIMATOR__1mRUQ_FwFtU8SxZ8-7fcfmPczAj6o3E4JHtpY_KQhqHk.xlsx'
WORK = SOURCE / 'work/cost'
PROFILE = WORK / 'lo-profile'
REPO = Path(__file__).resolve().parent.parent
FIXTURES = REPO / 'src/lib/cost/__tests__/fixtures'

IH = 'INSERT HERE'
QPI = 'QUANTITES PER ITEM'  # sic — the sheet's own spelling
OL = 'ORDER LIST'

# CostInputs key -> cell on INSERT HERE. Must match INPUT_CELLS in
# src/lib/cost/model.ts (the tests compare them). None = the sheet asks the
# question but has no answer cell.
INPUT_CELLS: dict[str, str | None] = {
    'longSideFt': 'B5',
    'shortSideFt': 'B7',
    'plantingSquares': 'B9',
    'naturePlaySquares': 'B11',
    'gravelEdgeFt': 'F19',
    'gravelEdgeOnHardscapeFt': 'F20',
    'gravelEdgeOnSoftscapeFt': 'F21',
    'outerEdgeFt': 'F23',
    'outerEdgeOnHardscapeFt': 'F24',
    'outerEdgeOnSoftscapeFt': 'F25',
    'outerEdgeGabionConnections': 'F26',
    'shrubs': 'C31',
    'smallTrees': 'C32',
    'largeTrees': 'C33',
    'gabionBaskets': 'B37',
    'raisedBedWoodEdgeFt': 'F44',
    'raisedBedGabionConnections': 'F45',
    'woodToppedGabions': None,  # question B51; answer cell deleted (F53 = #REF!)
    'benchesWithBackAndArms': 'B55',
    'benchesWithBack': 'B58',
    'benchesNoBack': 'B61',
    'squareTables': 'B64',
    'stools': 'B67',
    'gabionTables': 'B70',
    'stageSquares': 'B73',
    'trellises': 'B77',
    'longTables': 'B87',
    'compostBins': 'B90',
    'keyholeGardensLarge': 'B94',
    'keyholeGardensMedium': 'C94',
    'keyholeGardensSmall': 'D94',
    'sheds4x4': 'B104',
    'sheds4x8': 'D104',
    'cisterns4x4': 'B108',
    'cisterns4x8': 'D108',
    'rainBarrels': 'B111',
    'cafeTableSets': 'B114',
    'coldFrameSquares': 'B117',
    'optCafeTableSets': 'D126',
    'fountains': 'D127',
    'birdBaths': 'D128',
    'birdHouses': 'D129',
    'eventTents': 'D130',
    'adirondackChairs': 'D131',
    'hammocks': 'D132',
    'porchSwings': 'D133',
    'trashCans': 'D134',
    'solarLights': 'D135',
    'otherCosts': 'F145',
}

ZERO = {k: 0 for k in INPUT_CELLS}

# The values sitting in the orange cells of the downloaded sheet.
SHEET_DEFAULTS = {
    **ZERO,
    'longSideFt': 27, 'shortSideFt': 56, 'plantingSquares': 7, 'naturePlaySquares': 0,
    'shrubs': 10, 'smallTrees': 2, 'largeTrees': 0,
    'gabionTables': 4, 'trellises': 2, 'longTables': 5,
}

INPUT_SETS: list[tuple[str, str, dict]] = [
    ('sheet-defaults', 'The example values in the downloaded spreadsheet', SHEET_DEFAULTS),
    ('size-a-no-furnishings', 'Small size-A park (50 x 16), edges and plants only, zero furnishings', {
        **ZERO,
        'longSideFt': 50, 'shortSideFt': 16, 'plantingSquares': 6,
        'gravelEdgeFt': 24, 'gravelEdgeOnSoftscapeFt': 24,
        'outerEdgeFt': 40, 'outerEdgeOnHardscapeFt': 10, 'outerEdgeOnSoftscapeFt': 30, 'outerEdgeGabionConnections': 2,
        'shrubs': 4, 'smallTrees': 1,
    }),
    ('size-e-everything', 'Big size-E park (110 x 55) with every item, a 12-foot stage and other costs', {
        'longSideFt': 110, 'shortSideFt': 55, 'plantingSquares': 48, 'naturePlaySquares': 12,
        'gravelEdgeFt': 180, 'gravelEdgeOnHardscapeFt': 60, 'gravelEdgeOnSoftscapeFt': 120,
        'outerEdgeFt': 130, 'outerEdgeOnHardscapeFt': 30, 'outerEdgeOnSoftscapeFt': 100, 'outerEdgeGabionConnections': 4,
        'shrubs': 36, 'smallTrees': 8, 'largeTrees': 3,
        'gabionBaskets': 24, 'raisedBedWoodEdgeFt': 32, 'raisedBedGabionConnections': 6,
        'woodToppedGabions': 3, 'benchesWithBackAndArms': 2, 'benchesWithBack': 4, 'benchesNoBack': 3,
        'squareTables': 3, 'stools': 8, 'gabionTables': 2, 'stageSquares': 3, 'trellises': 3,
        'longTables': 2, 'compostBins': 2, 'keyholeGardensLarge': 1, 'keyholeGardensMedium': 1, 'keyholeGardensSmall': 1,
        'sheds4x4': 1, 'sheds4x8': 1, 'cisterns4x4': 1, 'cisterns4x8': 1, 'rainBarrels': 2, 'cafeTableSets': 2,
        'coldFrameSquares': 4,
        'optCafeTableSets': 2, 'fountains': 1, 'birdBaths': 2, 'birdHouses': 3, 'eventTents': 1,
        'adirondackChairs': 4, 'hammocks': 1, 'porchSwings': 2, 'trashCans': 2, 'solarLights': 20,
        'otherCosts': 500,
    }),
    ('size-c-fractions', 'Size-C park with fractional feet, an 8-foot stage, an odd number of cold-frame squares', {
        **ZERO,
        'longSideFt': 82.5, 'shortSideFt': 33.25, 'plantingSquares': 21, 'naturePlaySquares': 5,
        'gravelEdgeFt': 37, 'gravelEdgeOnHardscapeFt': 12, 'gravelEdgeOnSoftscapeFt': 25,
        'outerEdgeFt': 61.5, 'outerEdgeOnHardscapeFt': 20, 'outerEdgeOnSoftscapeFt': 41.5, 'outerEdgeGabionConnections': 3,
        'shrubs': 15, 'smallTrees': 3, 'largeTrees': 1,
        'gabionBaskets': 9, 'raisedBedGabionConnections': 3,
        'benchesWithBackAndArms': 1, 'benchesNoBack': 2, 'squareTables': 1, 'stools': 3, 'gabionTables': 1,
        'stageSquares': 2, 'trellises': 1, 'compostBins': 1, 'keyholeGardensMedium': 2,
        'sheds4x8': 1, 'rainBarrels': 1, 'coldFrameSquares': 3,
        'solarLights': 16, 'trashCans': 1, 'porchSwings': 1,
    }),
    ('size-b-stage16', 'Size-B park (64 x 28) with a 16-foot stage, benches with backs and fountains', {
        **ZERO,
        'longSideFt': 64, 'shortSideFt': 28, 'plantingSquares': 10, 'naturePlaySquares': 3,
        'shrubs': 6, 'smallTrees': 2,
        'gabionBaskets': 4, 'stageSquares': 4, 'benchesWithBack': 2,
        'coldFrameSquares': 2, 'fountains': 2, 'solarLights': 17, 'otherCosts': 125.5,
    }),
    ('rounding-boundaries', 'Values that land exactly on whole units (60 x 40, 75 planting squares) and a 1-square stage', {
        **ZERO,
        'longSideFt': 60, 'shortSideFt': 40, 'plantingSquares': 75,
        'gravelEdgeFt': 60, 'gravelEdgeOnHardscapeFt': 60,
        'outerEdgeFt': 48, 'outerEdgeOnSoftscapeFt': 48,
        'gabionBaskets': 27, 'stageSquares': 1, 'coldFrameSquares': 1, 'solarLights': 32,
        'hammocks': 2,
    }),
]

# Order-list rows that carry an item (ORDER LIST sheet).
ORDER_ROWS = [5, 6, 7, 10, 11, 15, 16, 17, 18, 20, 21, 22, 23, 29, 30, 31, 32,
              34, 35, 36, 37, 38, 39, 40, 41, 42, 46, 47, 48, 49, 50, 51, 52, 53, 54, 55,
              59, 60, 61, 62, 63, 64, 65, 66, 67, 68]

IH_OUTPUT_CELLS = ['C30', 'D30', 'D31', 'D32', 'D33', 'F30', 'F31', 'F32', 'F33', 'D40', 'D41', 'F40', 'F41',
                   'H16', 'H48', 'H80', 'H97', 'I22'] + [f'F{r}' for r in range(14, 147)]


def sheet_dump() -> None:
    wb = openpyxl.load_workbook(XLSX)
    wbv = openpyxl.load_workbook(XLSX, data_only=True)
    WORK.mkdir(parents=True, exist_ok=True)
    for ws in wb.worksheets:
        wv = wbv[ws.title]
        out = WORK / f'dump_{ws.title.replace(" ", "_")}.txt'
        with out.open('w') as f:
            for row in ws.iter_rows():
                for c in row:
                    if c.value is None:
                        continue
                    fill = c.fill.fgColor.rgb if c.fill and c.fill.fill_type else ''
                    if not isinstance(fill, str):
                        fill = ''
                    if isinstance(c.value, str) and c.value.startswith('='):
                        f.write(f'{c.coordinate}\t{c.value}\t=> {wv[c.coordinate].value!r}\t{fill}\n')
                    else:
                        f.write(f'{c.coordinate}\t{c.value!r}\t{fill}\n')
        print('wrote', out)


def ensure_profile() -> None:
    user = PROFILE / 'user'
    user.mkdir(parents=True, exist_ok=True)
    (user / 'registrymodifications.xcu').write_text(
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" '
        'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">\n'
        '<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>\n'
        '<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="ODFRecalcMode" oor:op="fuse"><value>0</value></prop></item>\n'
        '</oor:items>\n'
    )


def recalc(paths: list[Path], outdir: Path) -> None:
    ensure_profile()
    outdir.mkdir(parents=True, exist_ok=True)
    cmd = ['soffice', f'-env:UserInstallation=file://{PROFILE}', '--headless',
           '--convert-to', 'xlsx', '--outdir', str(outdir), *map(str, paths)]
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
    for p in paths:
        if not (outdir / p.name).exists():
            sys.exit(f'LibreOffice did not produce {p.name}:\n{r.stdout}\n{r.stderr}')


def num(v):
    """JSON-friendly cell value: numbers stay numbers, errors and text stay strings, blanks -> None."""
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return float(v)
    return v


def verify() -> None:
    """Recalculate the untouched file and compare with Google's cached values."""
    tmp = WORK / 'verify'
    if tmp.exists():
        shutil.rmtree(tmp)
    (tmp / 'in').mkdir(parents=True)
    # openpyxl round-trip drops the cached values, so LibreOffice has to compute everything.
    openpyxl.load_workbook(XLSX).save(tmp / 'in/original.xlsx')
    recalc([tmp / 'in/original.xlsx'], tmp / 'out')
    google = openpyxl.load_workbook(XLSX, data_only=True)
    lo = openpyxl.load_workbook(tmp / 'out/original.xlsx', data_only=True)
    formulas = openpyxl.load_workbook(XLSX)
    n = bad = 0
    for ws in formulas.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if not (isinstance(c.value, str) and c.value.startswith('=')):
                    continue
                n += 1
                a, b = google[ws.title][c.coordinate].value, lo[ws.title][c.coordinate].value
                if isinstance(a, (int, float)) and isinstance(b, (int, float)):
                    ok = abs(a - b) <= 1e-6 * max(1, abs(a))  # Google caches ~10 significant digits
                else:
                    ok = a == b or (a in (None, '') and b in (None, '', 0))
                if not ok:
                    bad += 1
                    print(f'  MISMATCH {ws.title}!{c.coordinate} {c.value[:60]} google={a!r} libreoffice={b!r}')
    print(f'verify: {n} formula cells, {bad} differ between Google Sheets and LibreOffice')
    if bad:
        sys.exit(1)


def fixtures() -> None:
    tmp = WORK / 'fixtures'
    if tmp.exists():
        shutil.rmtree(tmp)
    (tmp / 'in').mkdir(parents=True)
    for name, _, inputs in INPUT_SETS:
        assert set(inputs) == set(INPUT_CELLS), f'{name}: missing {set(INPUT_CELLS) - set(inputs)}'
        wb = openpyxl.load_workbook(XLSX)
        ws = wb[IH]
        for key, cell in INPUT_CELLS.items():
            if cell:
                ws[cell] = inputs[key]
        wb.save(tmp / 'in' / f'{name}.xlsx')
    recalc([tmp / 'in' / f'{n}.xlsx' for n, _, _ in INPUT_SETS], tmp / 'out')

    version = subprocess.run(['soffice', '--version'], capture_output=True, text=True).stdout.strip()
    FIXTURES.mkdir(parents=True, exist_ok=True)
    for name, description, inputs in INPUT_SETS:
        wb = openpyxl.load_workbook(tmp / 'out' / f'{name}.xlsx', data_only=True)
        ih, qpi, ol = wb[IH], wb[QPI], wb[OL]

        # Proof that LibreOffice recalculated with OUR inputs.
        L, W = inputs['longSideFt'], inputs['shortSideFt']
        assert abs(qpi['C5'].value - L * W) < 1e-9, (name, 'area not recalculated')
        assert abs(qpi['C12'].value - (2 * L + 2 * W)) < 1e-9, (name, 'perimeter not recalculated')

        insert_here = {c: num(ih[c].value) for c in IH_OUTPUT_CELLS if ih[c].value is not None}
        qpi_vals = {}
        for row in qpi.iter_rows(min_row=3, max_row=300, min_col=2, max_col=13):
            for c in row:
                if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                    qpi_vals[c.coordinate] = float(c.value)
        order = []
        for r in ORDER_ROWS:
            rec = {'row': r}
            for col in 'ACDEFGILMNO':
                v = ol[f'{col}{r}'].value
                if v is not None:
                    rec[col] = num(v)
            order.append(rec)
        totals = {c: num(ol[c].value) for c in ('P18', 'P23', 'O71')}

        out = {
            'name': name,
            'description': description,
            'source': 'PiaT cost-estimator spreadsheet (Dream workbook p.18), recalculated with the inputs below',
            'generatedBy': f'scripts/analyze_cost_model.py fixtures ({version})',
            'inputs': inputs,
            'cells': INPUT_CELLS,
            'insertHere': insert_here,
            'qpi': qpi_vals,
            'orderList': order,
            'orderListTotals': totals,
        }
        path = FIXTURES / f'{name}.json'
        path.write_text(json.dumps(out, indent=1, sort_keys=False) + '\n')
        print(f'wrote {path.relative_to(REPO)}  final cost {insert_here.get("F146")}  order list {totals["O71"]}')


if __name__ == '__main__':
    cmds = sys.argv[1:] or ['verify', 'fixtures']
    for cmd in cmds:
        {'dump': sheet_dump, 'verify': verify, 'fixtures': fixtures}[cmd]()
